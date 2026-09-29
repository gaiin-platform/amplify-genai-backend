import io
import json
import os

import pytest
from pptx import Presentation

from presentation_agent import pipeline
from presentation_agent.builder import build_deck
from presentation_agent.catalog import TemplateCatalog, build_catalog
from presentation_agent.plan import Bullet, DeckPlan, ReviewResult, SlideIssue, SlideRevision, SlideSpec
from presentation_agent.stores import ConsoleJobStore, LocalBlobStore, catalog_key
from presentation_agent.validate import validate_deck

HERE = os.path.dirname(__file__)
TEMPLATE_DIR = os.path.abspath(os.path.join(HERE, "..", "..", "..", "misc_deployment_files", "templates"))
TEMPLATES = sorted(f for f in os.listdir(TEMPLATE_DIR) if f.endswith(".pptx"))


def template_bytes(name: str) -> bytes:
    with open(os.path.join(TEMPLATE_DIR, name), "rb") as f:
        return f.read()


@pytest.fixture
def plan() -> DeckPlan:
    with open(os.path.join(HERE, "fixtures", "sample_plan.json")) as f:
        return DeckPlan.model_validate(json.load(f))


@pytest.mark.parametrize("name", TEMPLATES)
def test_catalog_finds_core_layouts(name):
    catalog = build_catalog(template_bytes(name))
    kinds = {l.kind for l in catalog.layouts}
    assert {"title", "title_content", "two_content", "title_only"} <= kinds
    assert all("vertical" not in l.name.lower() for l in catalog.layouts)
    comparison = catalog.first_of_kind("comparison")
    if comparison:
        assert len(comparison.regions_by_role("heading")) == 2
        assert len(comparison.regions_by_role("body")) == 2


def test_catalog_round_trips_through_json():
    catalog = build_catalog(template_bytes("celestial.pptx"))
    restored = TemplateCatalog.from_dict(json.loads(json.dumps(catalog.to_dict())))
    assert restored.prompt_summary() == catalog.prompt_summary()


@pytest.mark.parametrize("name", TEMPLATES)
def test_build_deck_places_every_slide(name, plan):
    template = template_bytes(name)
    catalog = build_catalog(template)
    deck = Presentation(io.BytesIO(build_deck(template, catalog, plan)))

    assert len(deck.slides) == len(plan.slides)  # template example slides removed
    for slide, spec in zip(deck.slides, plan.slides):
        texts = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        assert spec.title in texts
        assert "Click to" not in texts
        if spec.notes:
            assert slide.notes_slide.notes_text_frame.text == spec.notes
    kinds = [[s.shape_type for s in slide.shapes] for slide in deck.slides]
    assert any(s.has_chart for s in deck.slides[3].shapes)
    assert any(s.has_table for s in deck.slides[7].shapes)
    assert len(kinds[6]) > 4  # process diagram shapes


def test_invalid_layout_index_falls_back(plan):
    template = template_bytes("celestial.pptx")
    catalog = build_catalog(template)
    plan.slides = [SlideSpec(function="content", layout=999, title="Fallback works", bullets=[Bullet(text="a")])]
    deck = Presentation(io.BytesIO(build_deck(template, catalog, plan)))
    assert deck.slides[0].slide_layout.name == "Title and Content"


def test_validate_flags_overflow_and_placeholder_text(plan):
    template = template_bytes("celestial.pptx")
    catalog = build_catalog(template)
    plan.slides = [
        SlideSpec(
            function="content",
            layout=1,
            title="A title",
            bullets=[Bullet(text="Lorem ipsum " + "word " * 40) for _ in range(12)],
        )
    ]
    kinds = {w.kind for w in validate_deck(build_deck(template, catalog, plan), catalog)}
    assert {"text_overflow", "placeholder_text", "dense_text", "too_many_bullets"} <= kinds


def test_pipeline_applies_review_revisions(tmp_path, monkeypatch, plan):
    blobs = LocalBlobStore(str(tmp_path), template_dir=TEMPLATE_DIR)
    fixed = plan.slides[1].model_copy(update={"title": "Revised agenda title"})
    responses = [
        plan,
        ReviewResult(score=6, issues=[SlideIssue(slide=2, severity="medium", problem="weak")], revisions=[SlideRevision(slide=2, spec=fixed)]),
        ReviewResult(score=9),
    ]
    calls = []

    def fake_structured(self, system, content, output_model):
        calls.append(output_model.__name__)
        self.usage.add(self.model_id, {"inputTokens": 100, "outputTokens": 10})
        return responses.pop(0)

    monkeypatch.setattr(pipeline.LLM, "structured", fake_structured)
    fake_png = _png()
    request = pipeline.GenerateRequest(job_id="job-1", user="tester", template_name="celestial.pptx", content="notes", max_review_passes=2)
    jobs = ConsoleJobStore()
    result = pipeline.generate(request, blobs, jobs, renderer=lambda deck: [fake_png] * len(plan.slides))

    assert calls == ["DeckPlan", "ReviewResult", "ReviewResult"]
    assert result.review_scores == [6, 9]
    assert result.slide_count == len(plan.slides)
    deck = Presentation(io.BytesIO(blobs.get(result.pptx_key)))
    assert "Revised agenda title" in deck.slides[1].shapes.title.text
    assert all(blobs.get(k) for k in result.slide_keys)
    assert sum(u["inputTokens"] for u in result.usage.values()) == 300


def test_pipeline_prefers_induced_catalog(tmp_path):
    blobs = LocalBlobStore(str(tmp_path), template_dir=TEMPLATE_DIR)
    catalog = build_catalog(template_bytes("celestial.pptx"))
    catalog.layouts[1].description = "Induced description"
    blobs.put_json(catalog_key("celestial.pptx"), catalog.to_dict())
    loaded = pipeline.load_catalog(blobs, "celestial.pptx", template_bytes("celestial.pptx"))
    assert loaded.layouts[1].description == "Induced description"


def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (160, 90), "white").save(buf, format="PNG")
    return buf.getvalue()
