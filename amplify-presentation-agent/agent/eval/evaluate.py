"""Score decks on PPTEval's three dimensions (content, design, coherence; 1-5).

Adapted from icip-cas/PPTAgent's PPTEval (MIT): a vision model judges each
rendered slide for content and design, and the whole deck's text for
coherence. Use it to compare the agent against the pandoc baseline on a fixed
prompt set:

  # score existing decks (e.g. pandoc exports downloaded from the old UI)
  python -m eval.evaluate score --decks ./baseline/*.pptx --out baseline.csv

  # generate with the agent from prompts/*.md, then score
  python -m eval.evaluate generate --prompts ./eval/prompts --templates ../../misc_deployment_files/templates \
      --template celestial.pptx --workdir ./eval-out --out agent.csv
"""

import argparse
import csv
import glob
import io
import os
import statistics
import uuid
from typing import Dict, List

from pptx import Presentation
from pydantic import BaseModel, Field

from presentation_agent.llm import LLM, UsageTracker, image_block, text_block
from presentation_agent.pipeline import GenerateRequest, generate
from presentation_agent.render import render_slides
from presentation_agent.stores import ConsoleJobStore, LocalBlobStore

SLIDE_JUDGE = """You are grading a single presentation slide image on two 1-5 scales.
Content (1-5): 5 = one clear, specific message with well-edited supporting points or a fitting visual; 3 = understandable but generic or wordy; 1 = empty, confusing, or placeholder text.
Design (1-5): 5 = clean, legible, balanced layout with no overflow/overlap and good use of the visual hierarchy; 3 = acceptable but cluttered or plain bullets only; 1 = broken layout, clipped or illegible text.
Be strict and consistent. Give a one-sentence reason."""

DECK_JUDGE = """You are grading the coherence of a whole presentation from its slide text, 1-5.
5 = clear narrative arc (context, evidence, implications, conclusion), logical order, no redundancy, strong opening and close; 3 = reasonable topical grouping but weak flow; 1 = disjointed or repetitive. Give a one-sentence reason."""


class SlideScore(BaseModel):
    content: int = Field(ge=1, le=5)
    design: int = Field(ge=1, le=5)
    reason: str


class DeckScore(BaseModel):
    coherence: int = Field(ge=1, le=5)
    reason: str


def deck_text(pptx_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(pptx_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        slides.append(f"Slide {i}:\n" + "\n".join(texts))
    return "\n\n".join(slides)


def score_deck(pptx_bytes: bytes, judge: LLM) -> Dict:
    slides = render_slides(pptx_bytes)
    slide_scores = [judge.structured(SLIDE_JUDGE, [text_block("Grade this slide."), image_block(png)], SlideScore) for png in slides]
    coherence = judge.structured(DECK_JUDGE, deck_text(pptx_bytes), DeckScore)
    return {
        "slides": len(slides),
        "content": round(statistics.mean(s.content for s in slide_scores), 2),
        "design": round(statistics.mean(s.design for s in slide_scores), 2),
        "coherence": coherence.coherence,
    }


def write_rows(rows: List[Dict], out: str) -> None:
    with open(out, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    for key in ("content", "design", "coherence"):
        print(f"mean {key}: {statistics.mean(r[key] for r in rows):.2f}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["score", "generate"])
    parser.add_argument("--decks", nargs="*", default=[])
    parser.add_argument("--prompts")
    parser.add_argument("--templates")
    parser.add_argument("--template")
    parser.add_argument("--workdir", default="./eval-out")
    parser.add_argument("--judge-model", help="Defaults to the agent's default model")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    judge = LLM(args.judge_model, UsageTracker())
    rows = []
    if args.command == "score":
        for path in args.decks:
            with open(path, "rb") as f:
                rows.append({"deck": os.path.basename(path), **score_deck(f.read(), judge)})
    else:
        blobs = LocalBlobStore(args.workdir, template_dir=args.templates)
        for prompt_path in sorted(glob.glob(os.path.join(args.prompts, "*.md"))):
            with open(prompt_path) as f:
                content = f.read()
            request = GenerateRequest(job_id=f"eval-{uuid.uuid4().hex[:8]}", user="eval", template_name=args.template, content=content)
            result = generate(request, blobs, ConsoleJobStore())
            deck = blobs.get(result.pptx_key)
            rows.append({"deck": os.path.basename(prompt_path), **score_deck(deck, judge), "reviewScores": result.review_scores})
    write_rows(rows, args.out)


if __name__ == "__main__":
    main()
