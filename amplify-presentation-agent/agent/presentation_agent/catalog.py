"""Deterministic analysis of a .pptx template.

Produces a layout catalog (which layouts exist, what regions they offer, how
much text each region holds) plus theme tokens. The planner picks layouts from
this catalog and the builder uses it to place content. `induct.py` can enrich
the catalog with model-written layout descriptions, but generation never
depends on that step.
"""

import io
from dataclasses import asdict, dataclass, field
from typing import Dict, List, Optional

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

EMU_PER_INCH = 914400
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
FURNITURE_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
CONTENT_TYPES = {
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.CHART,
    PP_PLACEHOLDER.TABLE,
    PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.MEDIA_CLIP,
    PP_PLACEHOLDER.ORG_CHART,
    PP_PLACEHOLDER.BITMAP,
}

# A BODY placeholder shorter than this is a heading/caption strip, not a content area.
SMALL_REGION_HEIGHT_IN = 1.2

DEFAULT_TITLE_PT = 36.0
DEFAULT_BODY_PT = 20.0


@dataclass
class Region:
    idx: int
    role: str  # title | subtitle | body | caption | heading | picture
    left: float
    top: float
    width: float
    height: float
    font_pt: float
    max_lines: int
    chars_per_line: int


@dataclass
class LayoutInfo:
    index: int
    name: str
    kind: str
    regions: List[Region]
    description: str = ""
    best_for: List[str] = field(default_factory=list)

    def region(self, role: str) -> Optional[Region]:
        return next((r for r in self.regions if r.role == role), None)

    def regions_by_role(self, role: str) -> List[Region]:
        return sorted((r for r in self.regions if r.role == role), key=lambda r: (r.left, r.top))


@dataclass
class TemplateCatalog:
    slide_width: float
    slide_height: float
    heading_font: Optional[str]
    body_font: Optional[str]
    accent_colors: List[str]
    layouts: List[LayoutInfo]
    example_slide_count: int
    enriched: bool = False

    def layout(self, index: int) -> Optional[LayoutInfo]:
        return next((l for l in self.layouts if l.index == index), None)

    def first_of_kind(self, *kinds: str) -> Optional[LayoutInfo]:
        for kind in kinds:
            for layout in self.layouts:
                if layout.kind == kind:
                    return layout
        return None

    def to_dict(self) -> Dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: Dict) -> "TemplateCatalog":
        layouts = [
            LayoutInfo(
                index=l["index"],
                name=l["name"],
                kind=l["kind"],
                regions=[Region(**r) for r in l["regions"]],
                description=l.get("description", ""),
                best_for=l.get("best_for", []),
            )
            for l in data["layouts"]
        ]
        return cls(
            slide_width=data["slide_width"],
            slide_height=data["slide_height"],
            heading_font=data.get("heading_font"),
            body_font=data.get("body_font"),
            accent_colors=data.get("accent_colors", []),
            layouts=layouts,
            example_slide_count=data.get("example_slide_count", 0),
            enriched=data.get("enriched", False),
        )

    def prompt_summary(self) -> str:
        """Compact description of usable layouts for the planner."""
        lines = [
            f'Slide size: {self.slide_width:.2f}" x {self.slide_height:.2f}". '
            f"Heading font: {self.heading_font or 'theme default'}. Body font: {self.body_font or 'theme default'}.",
            "Layouts (use the index in `layout`):",
        ]
        for layout in self.layouts:
            regions = ", ".join(
                f"{r.role} ~{r.max_lines} lines x {r.chars_per_line} chars @{r.font_pt:.0f}pt"
                for r in layout.regions
                if r.role != "title"
            ) or "title only"
            desc = f" - {layout.description}" if layout.description else ""
            best = f" Best for: {', '.join(layout.best_for)}." if layout.best_for else ""
            lines.append(f"- [{layout.index}] {layout.name} (kind: {layout.kind}){desc}. Regions: {regions}.{best}")
        return "\n".join(lines)


def _inches(value: Optional[int]) -> float:
    return round((value or 0) / EMU_PER_INCH, 2)


def _sz_from_lst_style(element, level_tag: str = "a:lvl1pPr") -> Optional[float]:
    if element is None:
        return None
    found = element.find(f".//{level_tag}/a:defRPr", NS)
    if found is not None and found.get("sz"):
        return int(found.get("sz")) / 100.0
    return None


def _master_text_size(master, style: str) -> Optional[float]:
    tx_styles = master.element.find("p:txStyles", NS)
    if tx_styles is None:
        return None
    return _sz_from_lst_style(tx_styles.find(f"p:{style}", NS))


def _placeholder_font_pt(placeholder, master, is_title: bool) -> float:
    size = _sz_from_lst_style(placeholder.element.find(".//a:lstStyle", NS)) if placeholder.has_text_frame else None
    if size is None:
        ph_type = placeholder.placeholder_format.type
        for master_ph in master.placeholders:
            master_type = master_ph.placeholder_format.type
            same = master_type == ph_type or (is_title and master_type in TITLE_TYPES)
            if same:
                size = _sz_from_lst_style(master_ph.element.find(".//a:lstStyle", NS))
                break
    if size is None:
        size = _master_text_size(master, "titleStyle" if is_title else "bodyStyle")
    return size or (DEFAULT_TITLE_PT if is_title else DEFAULT_BODY_PT)


def _capacity(width_in: float, height_in: float, font_pt: float):
    # Average glyph width ~0.5em and line height ~1.2em, less text-frame insets.
    usable_w = max(width_in - 0.2, 0.1) * 72
    usable_h = max(height_in - 0.1, 0.1) * 72
    chars = max(int(usable_w / (font_pt * 0.5)), 1)
    lines = max(int(usable_h / (font_pt * 1.2)), 1)
    return lines, chars


def _overlap_x(a: Region, b: Region) -> float:
    """Fraction of the narrower region's width that the two regions share horizontally."""
    shared = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
    return max(shared, 0) / max(min(a.width, b.width), 0.01)


def _classify(name: str, regions: List[Region]) -> str:
    lowered = name.lower()
    roles = [r.role for r in regions]
    bodies = roles.count("body")
    headings = roles.count("heading")
    has_title = "title" in roles
    if "vertical" in lowered:
        return "unsupported"
    if "subtitle" in roles:
        return "title"
    if "picture" in roles:
        return "picture_caption"
    if not has_title and not regions:
        return "blank"
    if has_title and len(roles) == 1:
        return "title_only"
    if ("section" in lowered or "header" in lowered) and bodies <= 1 and headings == 0:
        return "section"
    if headings == 3 and bodies == 3:
        return "three_column"
    if headings == 2 and bodies == 2:
        return "comparison"
    if bodies == 2:
        return "two_content"
    if bodies == 1 and "caption" in roles:
        return "content_caption"
    if bodies == 1:
        return "title_content"
    if "quote" in lowered:
        return "quote"
    if has_title and "caption" in roles:
        return "title_caption"
    return "other"


def _theme_tokens(prs):
    heading_font = body_font = None
    accents: List[str] = []
    try:
        theme_part = prs.slide_master.part.part_related_by(RT.THEME)
        theme = etree.fromstring(theme_part.blob)
        major = theme.find(".//a:majorFont/a:latin", NS)
        minor = theme.find(".//a:minorFont/a:latin", NS)
        heading_font = major.get("typeface") if major is not None else None
        body_font = minor.get("typeface") if minor is not None else None
        scheme = theme.find(".//a:clrScheme", NS)
        if scheme is not None:
            for i in range(1, 7):
                accent = scheme.find(f"a:accent{i}", NS)
                if accent is not None and len(accent):
                    color = accent[0].get("val") or accent[0].get("lastClr")
                    if color:
                        accents.append(color)
    except Exception:
        pass
    return heading_font, body_font, accents


def build_catalog(template_bytes: bytes) -> TemplateCatalog:
    prs = Presentation(io.BytesIO(template_bytes))
    master = prs.slide_master
    layouts: List[LayoutInfo] = []

    for index, layout in enumerate(prs.slide_layouts):
        regions: List[Region] = []
        for ph in layout.placeholders:
            ph_type = ph.placeholder_format.type
            if ph_type in FURNITURE_TYPES:
                continue
            if ph.width is None or ph.height is None:
                continue
            width, height = _inches(ph.width), _inches(ph.height)
            is_title = ph_type in TITLE_TYPES
            if is_title:
                role = "title"
            elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                role = "subtitle"
            elif ph_type in (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BITMAP):
                role = "picture"
            elif ph_type in CONTENT_TYPES:
                role = "body" if height >= SMALL_REGION_HEIGHT_IN else "small"
            else:
                continue
            font_pt = _placeholder_font_pt(ph, master, is_title)
            lines, chars = _capacity(width, height, font_pt)
            regions.append(
                Region(
                    idx=ph.placeholder_format.idx,
                    role=role,
                    left=_inches(ph.left),
                    top=_inches(ph.top),
                    width=width,
                    height=height,
                    font_pt=font_pt,
                    max_lines=lines,
                    chars_per_line=chars,
                )
            )

        # Small text strips sitting directly above a content body are column headings
        # (Comparison layouts); any other small strip is a caption.
        bodies = [r for r in regions if r.role == "body"]
        for region in regions:
            if region.role != "small":
                continue
            above_body = any(_overlap_x(b, region) > 0.5 and -0.1 <= b.top - (region.top + region.height) < 1.0 for b in bodies)
            region.role = "heading" if above_body else "caption"

        # "Content with Caption" layouts carry a second, smaller body that describes the first.
        if "caption" in layout.name.lower() and len(bodies) == 2:
            min(bodies, key=lambda r: r.width * r.height).role = "caption"

        kind = _classify(layout.name, regions)
        if kind == "unsupported":
            continue
        layouts.append(LayoutInfo(index=index, name=layout.name, kind=kind, regions=regions))

    heading_font, body_font, accents = _theme_tokens(prs)
    return TemplateCatalog(
        slide_width=_inches(prs.slide_width),
        slide_height=_inches(prs.slide_height),
        heading_font=heading_font,
        body_font=body_font,
        accent_colors=accents,
        layouts=layouts,
        example_slide_count=len(prs.slides),
    )
