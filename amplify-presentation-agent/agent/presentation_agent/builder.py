"""Build a .pptx from a template and a DeckPlan.

Deterministic and side-effect free: the same template + plan always yields the
same deck, so review passes only ever edit the plan. Text goes into the
template's own placeholders so fonts, colors and bullet styles stay on-brand;
charts, tables, diagrams and images are placed inside the placeholder regions
they replace.
"""

import io
from dataclasses import dataclass
from typing import Dict, List, Optional

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .catalog import LayoutInfo, Region, TemplateCatalog
from .plan import Bullet, ChartSpec, DeckPlan, DiagramSpec, SlideSpec, TableSpec

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}

# Preferred layout kinds per slide function, used when the plan names a layout
# that does not exist or cannot hold the requested content.
FALLBACK_KINDS = {
    "opening": ("title",),
    "closing": ("title", "section", "title_only"),
    "section": ("section", "title", "title_only"),
    "agenda": ("title_content",),
    "content": ("title_content", "two_content"),
    "comparison": ("comparison", "two_content", "title_content"),
    "data": ("title_only", "title_content"),
    "process": ("title_only", "title_content"),
    "quote": ("title_only", "title_content"),
    "big_number": ("title_only", "title_content"),
}

MARGIN_IN = 0.6
GAP_IN = 0.3


@dataclass
class Rect:
    left: float
    top: float
    width: float
    height: float

    @classmethod
    def of(cls, region: Region) -> "Rect":
        return cls(region.left, region.top, region.width, region.height)

    def split_h(self, left_fraction: float) -> tuple:
        left_w = (self.width - GAP_IN) * left_fraction
        return (
            Rect(self.left, self.top, left_w, self.height),
            Rect(self.left + left_w + GAP_IN, self.top, self.width - left_w - GAP_IN, self.height),
        )

    def emu(self):
        return Inches(self.left), Inches(self.top), Inches(self.width), Inches(self.height)


def resolve_layout(catalog: TemplateCatalog, spec: SlideSpec) -> LayoutInfo:
    layout = catalog.layout(spec.layout)
    if layout is not None and layout.kind != "blank":
        return layout
    for kind in FALLBACK_KINDS.get(spec.function, ("title_content",)):
        layout = catalog.first_of_kind(kind)
        if layout is not None:
            return layout
    return catalog.layouts[0]


def build_deck(
    template_bytes: bytes,
    catalog: TemplateCatalog,
    plan: DeckPlan,
    images: Optional[Dict[int, bytes]] = None,
) -> bytes:
    """Return .pptx bytes. `images` maps 0-based slide index to generated image bytes."""
    prs = Presentation(io.BytesIO(template_bytes))
    _remove_all_slides(prs)
    images = images or {}

    for index, spec in enumerate(plan.slides):
        layout = resolve_layout(catalog, spec)
        slide = prs.slides.add_slide(prs.slide_layouts[layout.index])
        _SlideWriter(prs, slide, layout, catalog, spec, images.get(index)).write()

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _remove_all_slides(prs) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


class _SlideWriter:
    def __init__(self, prs, slide, layout: LayoutInfo, catalog: TemplateCatalog, spec: SlideSpec, image: Optional[bytes]):
        self.prs = prs
        self.slide = slide
        self.layout = layout
        self.catalog = catalog
        self.spec = spec
        self.image = image
        self.used: set = set()

    # ---- placeholder helpers -------------------------------------------------

    def _placeholder(self, region: Optional[Region]):
        if region is None:
            return None
        for ph in self.slide.placeholders:
            if ph.placeholder_format.idx == region.idx:
                return ph
        return None

    def _set_text(self, region: Optional[Region], text: str) -> bool:
        ph = self._placeholder(region)
        if ph is None or not text:
            return False
        ph.text_frame.text = text
        self.used.add(region.idx)
        return True

    def _set_bullets(self, region: Optional[Region], bullets: List[Bullet], heading: Optional[str] = None) -> bool:
        ph = self._placeholder(region)
        if ph is None or not (bullets or heading):
            return False
        tf = ph.text_frame
        tf.clear()
        first = True
        if heading:
            para = tf.paragraphs[0]
            para.text = heading
            para.font.bold = True
            first = False
        for bullet in bullets:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.text = bullet.text
            para.level = bullet.level
        # Keep the template's own autofit settings: forcing shrink-to-fit makes
        # LibreOffice (and PowerPoint on edit) shrink text far below the brand size.
        # Short lists in large regions read as empty; scale them up instead.
        lines = sum(max(1, -(-len(b.text) // region.chars_per_line)) for b in bullets) + (1 if heading else 0)
        if lines * 2.5 <= region.max_lines and region.font_pt < 24:
            size = Pt(min(region.font_pt * 1.4, 28))
            for para in tf.paragraphs:
                para.font.size = size
        self.used.add(region.idx)
        return True

    def _claim_region(self, region: Region) -> Rect:
        """Remove a placeholder so a chart/table/diagram can occupy its area."""
        ph = self._placeholder(region)
        if ph is not None:
            ph._element.getparent().remove(ph._element)
        self.used.add(region.idx)
        return Rect.of(region)

    def _free_area(self) -> Rect:
        """Content area below the title for layouts without a body region."""
        title = self.layout.region("title")
        width = self.catalog.slide_width
        height = self.catalog.slide_height
        if title is None:
            return Rect(MARGIN_IN, MARGIN_IN, width - 2 * MARGIN_IN, height - 2 * MARGIN_IN)
        if title.height > height * 0.5:  # side title (e.g. title column on the left)
            left = title.left + title.width + GAP_IN
            return Rect(left, title.top, width - left - MARGIN_IN, title.height)
        top = title.top + title.height + 0.15
        return Rect(title.left, top, title.width, height - top - 0.75)

    # ---- main -----------------------------------------------------------------

    def write(self) -> None:
        spec = self.spec
        layout = self.layout
        self._set_text(layout.region("title"), spec.title)

        if layout.kind in ("title", "section") or spec.function in ("opening", "closing", "section"):
            self._write_title_like()
        else:
            self._write_content()

        if spec.source:
            self._add_source(spec.source)
        if spec.notes:
            self.slide.notes_slide.notes_text_frame.text = spec.notes
        self._remove_unused_placeholders()

    def _write_title_like(self) -> None:
        spec = self.spec
        secondary = self.layout.region("subtitle") or self.layout.region("caption") or self.layout.region("body")
        text = spec.subtitle or "\n".join(b.text for b in spec.bullets)
        self._set_text(secondary, text)

    def _write_content(self) -> None:
        spec = self.spec
        layout = self.layout
        bodies = layout.regions_by_role("body")
        headings = layout.regions_by_role("heading")
        captions = layout.regions_by_role("caption")
        pictures = layout.regions_by_role("picture")
        visual = self._visual_kind()

        # Multi-column text (comparison / two / three column layouts).
        if spec.columns and len(bodies) >= 2 and not visual:
            for i, column in enumerate(spec.columns[: len(bodies)]):
                if i < len(headings) and column.heading:
                    self._set_text(headings[i], column.heading)
                    self._set_bullets(bodies[i], column.bullets)
                else:
                    self._set_bullets(bodies[i], column.bullets, heading=column.heading)
            if captions and spec.subtitle:
                self._set_text(captions[0], spec.subtitle)
            return

        # Image into a picture placeholder, text beside it.
        if visual == "image" and pictures and self.image:
            ph = self._placeholder(pictures[0])
            if ph is not None:
                ph.insert_picture(io.BytesIO(self.image))
                self.used.add(pictures[0].idx)
            text_region = (bodies or captions or [None])[0]
            if spec.bullets:
                self._set_bullets(text_region, spec.bullets)
            elif spec.subtitle:
                self._set_text(text_region, spec.subtitle)
            return

        text_blocks = spec.bullets or [b for c in spec.columns for b in c.bullets]

        if visual:
            if len(bodies) >= 2:
                if text_blocks:
                    self._set_bullets(bodies[0], text_blocks)
                    area = self._claim_region(bodies[1])
                else:
                    area = self._merge_regions(bodies)
            elif len(bodies) == 1:
                area = self._claim_region(bodies[0])
                if text_blocks:
                    text_area, area = area.split_h(0.4)
                    self._add_bullets_box(text_area, text_blocks, bodies[0].font_pt)
            else:
                area = self._free_area()
                if text_blocks:
                    text_area, area = area.split_h(0.4)
                    self._add_bullets_box(text_area, text_blocks, self._body_pt())
            self._draw_visual(visual, area)
            caption = captions[0] if captions else None
            if caption and spec.subtitle:
                self._set_text(caption, spec.subtitle)
            return

        # Plain text slide.
        if bodies:
            self._set_bullets(bodies[0], text_blocks)
            if len(bodies) > 1 and spec.subtitle:
                self._set_text(bodies[1], spec.subtitle)
        else:
            area = self._free_area()
            if text_blocks:
                self._add_bullets_box(area, text_blocks, self._body_pt())
        if captions and spec.subtitle:
            self._set_text(captions[0], spec.subtitle)

    def _merge_regions(self, regions: List[Region]) -> Rect:
        rects = [self._claim_region(r) for r in regions]
        left = min(r.left for r in rects)
        top = min(r.top for r in rects)
        right = max(r.left + r.width for r in rects)
        bottom = max(r.top + r.height for r in rects)
        return Rect(left, top, right - left, bottom - top)

    def _visual_kind(self) -> Optional[str]:
        spec = self.spec
        if spec.chart:
            return "chart"
        if spec.table:
            return "table"
        if spec.diagram:
            return "diagram"
        if spec.big_number:
            return "big_number"
        if spec.quote:
            return "quote"
        if spec.image and self.image:
            return "image"
        return None

    def _body_pt(self) -> float:
        body = self.catalog.first_of_kind("title_content")
        region = body.region("body") if body else None
        return region.font_pt if region else 18.0

    # ---- visuals ---------------------------------------------------------------

    def _draw_visual(self, kind: str, area: Rect) -> None:
        spec = self.spec
        if kind == "chart":
            _add_chart(self.slide, spec.chart, area)
        elif kind == "table":
            _add_table(self.slide, spec.table, area)
        elif kind == "diagram":
            _add_diagram(self.slide, spec.diagram, area)
        elif kind == "big_number":
            self._add_big_number(area)
        elif kind == "quote":
            self._add_quote(area)
        elif kind == "image":
            _add_picture_fill(self.slide, self.image, area)

    def _add_bullets_box(self, area: Rect, bullets: List[Bullet], font_pt: float) -> None:
        box = self.slide.shapes.add_textbox(*area.emu())
        tf = box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = ("• " if bullet.level == 0 else "– ") + bullet.text
            para.level = bullet.level
            para.font.size = Pt(font_pt if bullet.level == 0 else font_pt - 2)
            para.space_after = Pt(6)

    def _add_big_number(self, area: Rect) -> None:
        big = self.spec.big_number
        box = self.slide.shapes.add_textbox(*area.emu())
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        value = tf.paragraphs[0]
        value.text = big.value
        value.alignment = PP_ALIGN.CENTER
        value.font.bold = True
        value.font.size = Pt(max(40, min(120, area.height * 72 * 0.45)))
        value.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        label = tf.add_paragraph()
        label.text = big.label
        label.alignment = PP_ALIGN.CENTER
        label.font.size = Pt(22)

    def _add_quote(self, area: Rect) -> None:
        quote = self.spec.quote
        box = self.slide.shapes.add_textbox(*area.emu())
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        text = tf.paragraphs[0]
        text.text = f"“{quote.text}”"
        text.font.italic = True
        text.font.size = Pt(28)
        if quote.attribution:
            who = tf.add_paragraph()
            who.text = f"— {quote.attribution}"
            who.alignment = PP_ALIGN.RIGHT
            who.font.size = Pt(18)

    def _add_source(self, source: str) -> None:
        height = self.catalog.slide_height
        box = self.slide.shapes.add_textbox(Inches(MARGIN_IN), Inches(height - 0.55), Inches(self.catalog.slide_width - 2 * MARGIN_IN - 1.5), Inches(0.3))
        para = box.text_frame.paragraphs[0]
        para.text = f"Source: {source}"
        para.font.size = Pt(10)
        para.font.italic = True
        box.name = "Source"

    def _remove_unused_placeholders(self) -> None:
        for ph in list(self.slide.placeholders):
            if ph.placeholder_format.idx in self.used:
                continue
            if ph.has_text_frame and ph.text_frame.text.strip():
                continue
            ph._element.getparent().remove(ph._element)


def _add_chart(slide, chart: ChartSpec, area: Rect) -> None:
    data = CategoryChartData()
    data.categories = chart.categories
    for series in chart.series:
        values = list(series.values)[: len(chart.categories)]
        values += [0] * (len(chart.categories) - len(values))
        data.add_series(series.name, values, number_format=chart.number_format)
    chart_type = CHART_TYPES[chart.type]
    graphic = slide.shapes.add_chart(chart_type, *area.emu(), data)
    c = graphic.chart
    c.font.size = Pt(14)
    single = len(chart.series) == 1
    is_round = chart.type in ("pie", "doughnut")
    c.has_legend = is_round or not single
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
    if chart.title:
        c.has_title = True
        c.chart_title.text_frame.text = chart.title
        c.chart_title.text_frame.paragraphs[0].font.size = Pt(16)
    else:
        c.has_title = False
    plot = c.plots[0]
    if single or is_round:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(12)
        if chart.number_format:
            labels.number_format = chart.number_format
            labels.number_format_is_linked = False
        if is_round:
            labels.show_percentage = chart.number_format is None
            labels.show_value = chart.number_format is not None
    if not is_round:
        c.value_axis.has_major_gridlines = False
        c.value_axis.visible = not single
        if hasattr(plot, "gap_width"):
            plot.gap_width = 80
        if single and hasattr(plot, "vary_by_categories"):
            plot.vary_by_categories = False


def _add_table(slide, table: TableSpec, area: Rect) -> None:
    cols = max(len(table.headers), max((len(r) for r in table.rows), default=0), 1)
    rows = len(table.rows) + 1
    font_pt = 14 if rows <= 6 else 12 if rows <= 10 else 10
    row_h = min(area.height / rows, 0.6)
    shape = slide.shapes.add_table(rows, cols, Inches(area.left), Inches(area.top), Inches(area.width), Inches(row_h * rows))
    grid = shape.table
    for c in range(cols):
        grid.columns[c].width = Inches(area.width / cols)
    for c, header in enumerate(table.headers[:cols]):
        _set_cell(grid.cell(0, c), header, font_pt, bold=True)
    for r, row in enumerate(table.rows, start=1):
        for c in range(cols):
            _set_cell(grid.cell(r, c), row[c] if c < len(row) else "", font_pt)


def _set_cell(cell, text: str, font_pt: float, bold: bool = False) -> None:
    cell.text = text
    para = cell.text_frame.paragraphs[0]
    para.font.size = Pt(font_pt)
    para.font.bold = bold
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _add_diagram(slide, diagram: DiagramSpec, area: Rect) -> None:
    if diagram.kind == "timeline":
        _add_timeline(slide, diagram, area)
    else:
        _add_process(slide, diagram, area)


def _add_process(slide, diagram: DiagramSpec, area: Rect) -> None:
    steps = diagram.steps
    n = len(steps)
    arrow_w = 0.35
    gap = 0.12
    box_w = (area.width - (n - 1) * (arrow_w + 2 * gap)) / n
    has_detail = any(s.detail for s in steps)
    box_h = min(1.3, area.height * (0.4 if has_detail else 0.7))
    box_top = area.top + (area.height * 0.15 if has_detail else (area.height - box_h) / 2)
    label_pt = 18 if box_w > 2.2 else 16 if box_w > 1.6 else 14
    for i, step in enumerate(steps):
        left = area.left + i * (box_w + arrow_w + 2 * gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(box_top), Inches(box_w), Inches(box_h))
        box.fill.solid()
        box.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.text = step.label
        para.alignment = PP_ALIGN.CENTER
        para.font.bold = True
        para.font.size = Pt(label_pt)
        para.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
        if step.detail:
            detail = slide.shapes.add_textbox(Inches(left), Inches(box_top + box_h + 0.15), Inches(box_w), Inches(area.height - box_h - 0.3 - (box_top - area.top)))
            dtf = detail.text_frame
            dtf.word_wrap = True
            dpara = dtf.paragraphs[0]
            dpara.text = step.detail
            dpara.alignment = PP_ALIGN.CENTER
            dpara.font.size = Pt(14)
        if i < n - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(left + box_w + gap),
                Inches(box_top + box_h / 2 - 0.18),
                Inches(arrow_w),
                Inches(0.36),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
            arrow.line.fill.background()


def _add_timeline(slide, diagram: DiagramSpec, area: Rect) -> None:
    steps = diagram.steps
    n = len(steps)
    mid = area.top + area.height * 0.4
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(area.left), Inches(mid), Inches(area.left + area.width), Inches(mid))
    line.line.width = Pt(3)
    line.line.color.theme_color = MSO_THEME_COLOR.ACCENT_2
    slot = area.width / n
    dot = 0.32
    for i, step in enumerate(steps):
        center = area.left + slot * (i + 0.5)
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center - dot / 2), Inches(mid - dot / 2), Inches(dot), Inches(dot))
        marker.fill.solid()
        marker.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        marker.line.fill.background()
        label = slide.shapes.add_textbox(Inches(center - slot / 2 + 0.05), Inches(mid - 0.95), Inches(slot - 0.1), Inches(0.7))
        ltf = label.text_frame
        ltf.word_wrap = True
        ltf.vertical_anchor = MSO_ANCHOR.BOTTOM
        lpara = ltf.paragraphs[0]
        lpara.text = step.label
        lpara.alignment = PP_ALIGN.CENTER
        lpara.font.bold = True
        lpara.font.size = Pt(16)
        if step.detail:
            detail = slide.shapes.add_textbox(Inches(center - slot / 2 + 0.05), Inches(mid + 0.3), Inches(slot - 0.1), Inches(area.top + area.height - mid - 0.3))
            dtf = detail.text_frame
            dtf.word_wrap = True
            dpara = dtf.paragraphs[0]
            dpara.text = step.detail
            dpara.alignment = PP_ALIGN.CENTER
            dpara.font.size = Pt(13)


def _add_picture_fill(slide, image: bytes, area: Rect) -> None:
    """Place an image covering `area`, cropping (never stretching) to its aspect ratio."""
    with Image.open(io.BytesIO(image)) as img:
        img_ratio = img.width / img.height
    area_ratio = area.width / area.height
    pic = slide.shapes.add_picture(io.BytesIO(image), *area.emu())
    if img_ratio > area_ratio:
        excess = 1 - area_ratio / img_ratio
        pic.crop_left = pic.crop_right = excess / 2
    elif img_ratio < area_ratio:
        excess = 1 - img_ratio / area_ratio
        pic.crop_top = pic.crop_bottom = excess / 2
