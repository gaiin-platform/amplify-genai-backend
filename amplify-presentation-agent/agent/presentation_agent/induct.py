"""One-time template induction, run when an admin uploads a template.

Adapted from icip-cas/PPTAgent's SlideInducter (MIT): instead of clustering
example slides with ViT embeddings, we render one labeled sample per layout
and ask a vision model to describe each layout pattern and the slide functions
it suits. The deterministic catalog stays the source of truth for geometry;
the model only adds descriptions, suitability, and filters broken layouts.
"""

import io
import logging
from typing import Callable, Dict, List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pydantic import BaseModel, Field

from .builder import _remove_all_slides
from .catalog import TemplateCatalog, build_catalog
from .llm import LLM, UsageTracker, image_block, text_block
from .prompts import INDUCT_SYSTEM
from .render import render_slides
from .stores import BlobStore, catalog_key

logger = logging.getLogger("presentation_agent.induct")

MAX_LAYOUT_IMAGES = 16


class LayoutInsight(BaseModel):
    index: int
    description: str
    best_for: List[str] = Field(default_factory=list)
    usable: bool = True


class TemplateInsights(BaseModel):
    layouts: List[LayoutInsight]


def labeled_samples(template: bytes, catalog: TemplateCatalog) -> bytes:
    """One slide per catalog layout with every region labeled by its role."""
    prs = Presentation(io.BytesIO(template))
    _remove_all_slides(prs)
    for layout in catalog.layouts:
        slide = prs.slides.add_slide(prs.slide_layouts[layout.index])
        roles = {r.idx: r.role for r in layout.regions}
        for ph in list(slide.placeholders):
            role = roles.get(ph.placeholder_format.idx)
            if role is None:
                continue
            label = f"{role.upper()} REGION"
            if role == "picture":
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ph.left, ph.top, ph.width, ph.height)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
                shape.text_frame.text = label
                ph._element.getparent().remove(ph._element)
            elif ph.has_text_frame:
                ph.text_frame.text = label if role == "title" else f"{label}\nSecond line of sample text"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def induct_template(
    template_name: str,
    blobs: BlobStore,
    model_id: str,
    renderer: Callable[[bytes], List[bytes]] = render_slides,
) -> Dict:
    template = blobs.get_template(template_name)
    catalog = build_catalog(template)
    usage = UsageTracker()

    samples = renderer(labeled_samples(template, catalog))
    blocks = [
        text_block(
            "Template layouts (deterministic analysis):\n" + catalog.prompt_summary() + "\n\nRendered samples follow, in the same order."
        )
    ]
    for layout, png in list(zip(catalog.layouts, samples))[:MAX_LAYOUT_IMAGES]:
        blocks.append(text_block(f"Layout [{layout.index}] {layout.name}"))
        blocks.append(image_block(png))

    insights = LLM(model_id, usage).structured(INDUCT_SYSTEM, blocks, TemplateInsights)
    by_index = {i.index: i for i in insights.layouts}
    kept = []
    for layout in catalog.layouts:
        insight = by_index.get(layout.index)
        if insight is None:
            kept.append(layout)
            continue
        if not insight.usable and layout.kind not in ("title", "title_content", "title_only"):
            logger.info("Dropping layout %d (%s) judged unusable", layout.index, layout.name)
            continue
        layout.description = insight.description
        layout.best_for = insight.best_for
        kept.append(layout)
    catalog.layouts = kept
    catalog.enriched = True
    blobs.put_json(catalog_key(template_name), catalog.to_dict())
    return {"catalog": catalog.to_dict(), "usage": usage.as_dict()}
