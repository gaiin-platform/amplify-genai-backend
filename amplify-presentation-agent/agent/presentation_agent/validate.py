"""Structural and geometric lint for a built deck.

Warnings are heuristics fed to the visual reviewer alongside the rendered
slides; the reviewer decides which are real. Checks mirror the structural gate
in the sample's qa-guide (bounds, overflow risk, leftover placeholder text,
density) but are implemented here on python-pptx.
"""

import io
import math
import re
from dataclasses import asdict, dataclass
from typing import Dict, List, Optional

from pptx import Presentation

from .catalog import EMU_PER_INCH, TemplateCatalog

LEFTOVER_PATTERNS = re.compile(r"click to (add|edit)|lorem ipsum|\bTODO\b|\bTBD\b|\[insert|xx+%|<[A-Z _]+>", re.IGNORECASE)
MAX_BODY_WORDS = 75
MAX_BULLETS = 7
BOUNDS_TOLERANCE_IN = 0.05


@dataclass
class Warning:
    slide: int  # 1-based
    kind: str
    detail: str

    def to_dict(self) -> Dict:
        return asdict(self)


def _inches(emu: Optional[int]) -> float:
    return (emu or 0) / EMU_PER_INCH


def _run_size(text_frame) -> Optional[float]:
    for para in text_frame.paragraphs:
        if para.font.size is not None:
            return para.font.size.pt
        for run in para.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return None


def _needed_lines(text_frame, width_in: float, font_pt: float) -> int:
    lines = 0
    for para in text_frame.paragraphs:
        indent = 0.35 * (para.level or 0)
        size = para.font.size.pt if para.font.size is not None else font_pt
        usable = max(width_in - 0.2 - indent, 0.3) * 72
        per_line = max(int(usable / (size * 0.5)), 1)
        text = "".join(run.text for run in para.runs) or para.text
        lines += max(1, math.ceil(len(text) / per_line)) if text.strip() else 1
    return lines


def validate_deck(pptx_bytes: bytes, catalog: TemplateCatalog) -> List[Warning]:
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide_w = _inches(prs.slide_width)
    slide_h = _inches(prs.slide_height)
    layout_index = {id(layout): i for i, layout in enumerate(prs.slide_layouts)}
    warnings: List[Warning] = []

    for number, slide in enumerate(prs.slides, start=1):
        layout = catalog.layout(layout_index.get(id(slide.slide_layout), -1))
        region_fonts = {r.idx: r.font_pt for r in layout.regions} if layout else {}
        has_title = False
        text_boxes = []

        for shape in slide.shapes:
            left, top = _inches(shape.left), _inches(shape.top)
            width, height = _inches(shape.width), _inches(shape.height)
            if (
                left < -BOUNDS_TOLERANCE_IN
                or top < -BOUNDS_TOLERANCE_IN
                or left + width > slide_w + BOUNDS_TOLERANCE_IN
                or top + height > slide_h + BOUNDS_TOLERANCE_IN
            ):
                warnings.append(Warning(number, "out_of_bounds", f"'{shape.name}' extends past the slide edge"))

            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            if not text.strip():
                continue
            is_placeholder = shape.is_placeholder
            idx = shape.placeholder_format.idx if is_placeholder else None
            if is_placeholder and idx == 0:
                has_title = True
            if LEFTOVER_PATTERNS.search(text):
                warnings.append(Warning(number, "placeholder_text", f"'{shape.name}' contains placeholder-like text: {text[:60]!r}"))

            font_pt = _run_size(shape.text_frame) or region_fonts.get(idx) or 18.0
            available = max(int((height - 0.1) * 72 / (font_pt * 1.2)), 1)
            needed = _needed_lines(shape.text_frame, width, font_pt)
            longest = max((len(w) for w in text.split()), default=0)
            per_line = max(int(max(width - 0.2, 0.3) * 72 / (font_pt * 0.6)), 1)  # wide glyphs: be conservative
            if longest > per_line:
                warnings.append(
                    Warning(number, "word_break", f"'{shape.name}' is too narrow for the word length at {font_pt:.0f}pt; words will split mid-word")
                )
            if needed > available:
                warnings.append(
                    Warning(number, "text_overflow", f"'{shape.name}' needs ~{needed} lines but fits ~{available} at {font_pt:.0f}pt")
                )
            if idx != 0:
                words = len(text.split())
                paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
                if words > MAX_BODY_WORDS:
                    warnings.append(Warning(number, "dense_text", f"'{shape.name}' has {words} words"))
                if len(paragraphs) > MAX_BULLETS:
                    warnings.append(Warning(number, "too_many_bullets", f"'{shape.name}' has {len(paragraphs)} bullets"))
            text_boxes.append((shape.name, left, top, width, height))

        if not has_title:
            warnings.append(Warning(number, "missing_title", "slide has no title text"))

        for i, a in enumerate(text_boxes):
            for b in text_boxes[i + 1 :]:
                overlap = _overlap_fraction(a[1:], b[1:])
                if overlap > 0.15:
                    warnings.append(Warning(number, "overlap", f"'{a[0]}' and '{b[0]}' overlap ({overlap:.0%})"))

    return warnings


def _overlap_fraction(a, b) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    w = min(ax + aw, bx + bw) - max(ax, bx)
    h = min(ay + ah, by + bh) - max(ay, by)
    if w <= 0 or h <= 0:
        return 0.0
    return (w * h) / max(min(aw * ah, bw * bh), 0.01)
